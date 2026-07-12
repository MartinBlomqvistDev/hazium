"""Adapter for EFSA's OpenFoodTox chemical hazards database.

OpenFoodTox is a structured IUCLID/OECD Harmonised Template export of EFSA's
toxicological knowledge base, published to Zenodo. Unlike the two KEMI
sources, it gives genuinely dated evidence: each scientific dossier carries
``LiteratureReference.DateOfEvaluation``, the date the underlying EFSA
opinion became public. This is the source that makes the temporal ``as_of``
cutoff meaningful rather than vacuous — a pre-2023 view can now contain real
EFSA conclusions, not just a frozen snapshot.

Three surfaces are extracted from the 18-sheet document graph:

* ``REF_SUB`` — substance identity (CAS, EC, PubChem). Used to enrich the
  identity spine with an authoritative source, not to invent identifiers.
* ``FLEX_SUM.Metabolites`` — parent-substance -> metabolite links, becoming
  ``DEGRADES_TO`` edges. This sheet carries no per-row date (a metabolic
  pathway is a property of the substance record, not a dated observation), so
  a link's ``known_at`` is the parent substance's *earliest dated EFSA
  assessment* (from ``DOSSIER``): a metabolite relationship reported in a 2008
  EFSA conclusion was knowable in 2008, not just at the 2026 export date. This
  is a real join over already-ingested dates, not an invented one. Only when
  the parent has no dated assessment does the link fall back to the export's
  own publication date, the same snapshot honesty as the KEMI register
  structure.
* ``DOSSIER`` — one row per EFSA scientific opinion/conclusion, dated and
  DOI-linked. Becomes ``SourceDocument`` facts with real ``known_at``.

The IUCLID export links tables through UUID ``Parent UUID`` chains rather
than foreign keys to a single subject. Only the chains actually needed are
resolved here: ``SUB -> REF_SUB`` for identity, and ``Metabolites`` /
``DOSSIER`` -> their subject ``SUB`` document for degradation and evidence.
Real data is imperfect: ~0.2% of CAS values are malformed (encoding
artifacts, placeholders); those substances fall back to a name-based id
rather than aborting ingestion.
"""

from __future__ import annotations

import json
import urllib.request
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl

from hazium.models import DegradationLink, SourceDocument, Substance
from hazium.resolve.ids import safe_substance_node_id

SOURCE = "efsa:openfoodtox"
PUBLISHER = "EFSA"
RECORD_ID = "19388272"
RECORD_API_URL = f"https://zenodo.org/api/records/{RECORD_ID}"
EXPORT_URL = (
    f"https://zenodo.org/api/records/{RECORD_ID}/files/OFT3.0%20export%20repository.xlsx/content"
)


@dataclass(frozen=True)
class _SubstanceIdentity:
    name: str
    cas_number: str | None
    ec_number: str | None
    pubchem_cid: int | None


@dataclass(frozen=True)
class OpenFoodToxIndex:
    """Cross-references parsed from one export, kept in memory.

    Loading the workbook is the expensive step; ``substances_from``,
    ``degradation_links_from``, and ``assessments_from`` are pure transforms
    over this index, independently testable without a real spreadsheet.
    """

    substances: dict[str, _SubstanceIdentity]  # SUB document uuid -> identity
    degradation_pairs: list[tuple[str, str]]  # (parent SUB uuid, metabolite SUB uuid)
    assessments: list[dict[str, Any]]  # raw dossier rows


def download_export(dest_dir: Path) -> Path:
    """Download the OpenFoodTox 3.0 export, skipping if already present."""
    dest_dir.mkdir(parents=True, exist_ok=True)
    dest = dest_dir / "openfoodtox_3.0.xlsx"
    if not dest.exists():
        urllib.request.urlretrieve(EXPORT_URL, dest)
    return dest


def record_publication_date() -> date:
    """The Zenodo record's own publication date, read from its metadata.

    Used as ``known_at`` for facts the export carries no per-row date for
    (degradation links): the export itself is dated, even where its rows
    are not.
    """
    request = urllib.request.Request(RECORD_API_URL, headers={"User-Agent": "hazium/0.0.1"})
    with urllib.request.urlopen(request) as response:
        record = json.load(response)
    return date.fromisoformat(record["metadata"]["publication_date"])


def load(xlsx_path: Path) -> OpenFoodToxIndex:
    """Parse the three cross-references this adapter needs, in one pass."""
    workbook = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        ref_sub = _index_ref_sub(workbook["REF_SUB"])
        substances = _index_substances(workbook["SUB"], ref_sub)
        degradation_pairs = _index_degradation_pairs(workbook["FLEX_SUM.Metabolites"])
        assessments = _index_assessments(workbook["DOSSIER"])
    finally:
        workbook.close()
    return OpenFoodToxIndex(
        substances=substances,
        degradation_pairs=degradation_pairs,
        assessments=assessments,
    )


def _rows(sheet: Any) -> tuple[dict[str, int], Any]:
    """A sheet's header-to-column index map and its data row iterator."""
    it = sheet.iter_rows(values_only=True)
    header = next(it)
    index = {name: i for i, name in enumerate(header) if name}
    return index, it


def _first_segment(value: Any) -> str | None:
    """IUCLID reference fields are slash-joined; the referenced doc is first."""
    return str(value).split("/")[0] if value else None


def _index_ref_sub(sheet: Any) -> dict[str, _SubstanceIdentity]:
    idx, rows = _rows(sheet)
    uuid_col = idx["Document UUID"]
    name_col = idx["ReferenceSubstanceName"]
    cas_col = idx["Inventory.CASNumber"]
    ec_col = idx["Inventory.InventoryEntry"]
    pubchem_col = idx["PUBCHEM CID"]
    out: dict[str, _SubstanceIdentity] = {}
    for row in rows:
        uuid = row[uuid_col]
        if not uuid or not row[name_col]:
            continue
        pubchem_raw = row[pubchem_col]
        ec_raw = row[ec_col]
        out[uuid] = _SubstanceIdentity(
            name=row[name_col],
            cas_number=str(row[cas_col]).strip() if row[cas_col] else None,
            ec_number=str(ec_raw).split("@")[0] if ec_raw else None,
            pubchem_cid=int(pubchem_raw) if pubchem_raw and str(pubchem_raw).isdigit() else None,
        )
    return out


def _index_substances(
    sheet: Any, ref_sub: dict[str, _SubstanceIdentity]
) -> dict[str, _SubstanceIdentity]:
    """SUB documents, identity resolved through their REF_SUB link.

    Keyed on the SUB document uuid (not REF_SUB), because that is what the
    metabolite and dossier sheets reference.
    """
    idx, rows = _rows(sheet)
    uuid_col = idx["Document UUID"]
    name_col = idx["ChemicalName"]
    ref_col = idx["ReferenceSubstance.ReferenceSubstance"]
    out: dict[str, _SubstanceIdentity] = {}
    for row in rows:
        uuid = row[uuid_col]
        if not uuid:
            continue
        ref_uuid = _first_segment(row[ref_col])
        identity = ref_sub.get(ref_uuid) if ref_uuid else None
        name = row[name_col] or (identity.name if identity else None)
        if not name:
            continue
        out[uuid] = _SubstanceIdentity(
            name=name,
            cas_number=identity.cas_number if identity else None,
            ec_number=identity.ec_number if identity else None,
            pubchem_cid=identity.pubchem_cid if identity else None,
        )
    return out


def _index_degradation_pairs(sheet: Any) -> list[tuple[str, str]]:
    idx, rows = _rows(sheet)
    parent_col = idx["Parent UUID"]
    metabolite_col = idx["ListMetabolites.Metabolites.LinkMetaboliteDataset"]
    pairs = []
    for row in rows:
        parent = _first_segment(row[parent_col])
        metabolite = row[metabolite_col]
        if parent and metabolite:
            pairs.append((parent, metabolite))
    return pairs


def _index_assessments(sheet: Any) -> list[dict[str, Any]]:
    idx, rows = _rows(sheet)
    subject_col = idx["DossierSubject.Name"]
    date_col = idx["LiteratureReference.DateOfEvaluation"]
    title_col = idx["LiteratureReference.EFSAOutputTitle"]
    doi_col = idx["LiteratureReference.LinkToPersistentIdentifier"]
    uuid_col = idx["Document UUID"]
    out = []
    for row in rows:
        subject = _first_segment(row[subject_col])
        published = _parse_date(row[date_col])
        if not subject or not published or not row[title_col]:
            continue
        out.append(
            {
                "dossier_uuid": row[uuid_col],
                "subject_uuid": subject,
                "published_at": published,
                "title": row[title_col],
                "doi": row[doi_col],
            }
        )
    return out


def _parse_date(value: Any) -> date | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value)[:10])
    except ValueError:
        return None


def _node_id(identity: _SubstanceIdentity) -> str:
    return safe_substance_node_id(cas_number=identity.cas_number, name=identity.name)


def substances_from(index: OpenFoodToxIndex, known_at: date) -> list[Substance]:
    """Every identified substance in the index, deduplicated by node id.

    ``Substance.cas_number`` reflects what the source actually claimed, even
    when malformed; only graph identity (via ``_node_id``) applies the safe
    fallback. Conflating the two would silently discard source data that a
    later, stricter entity-resolution pass might still be able to use.
    """
    seen: dict[str, Substance] = {}
    for identity in index.substances.values():
        node_id = _node_id(identity)
        if node_id in seen:
            continue
        seen[node_id] = Substance(
            name=identity.name,
            cas_number=identity.cas_number,
            ec_number=identity.ec_number,
            pubchem_cid=identity.pubchem_cid,
            source=SOURCE,
            known_at=known_at,
        )
    return list(seen.values())


def _earliest_assessment_by_substance(index: OpenFoodToxIndex) -> dict[str, date]:
    """Each substance's earliest dated EFSA assessment, keyed by node id.

    Used to back-date degradation links to when the metabolic relationship
    was actually knowable, rather than the export's snapshot date.
    """
    earliest: dict[str, date] = {}
    for row in index.assessments:
        subject = index.substances.get(row["subject_uuid"])
        if not subject:
            continue
        node_id = _node_id(subject)
        published: date = row["published_at"]
        if node_id not in earliest or published < earliest[node_id]:
            earliest[node_id] = published
    return earliest


def degradation_links_from(
    index: OpenFoodToxIndex, fallback_known_at: date
) -> list[DegradationLink]:
    """DEGRADES_TO facts: a substance's declared metabolic degradation.

    ``known_at`` is the parent substance's earliest dated EFSA assessment
    where one exists (a real join, not an invented date); ``fallback_known_at``
    (the export's publication date) applies only when the parent has no dated
    assessment at all. See the module docstring's dating note.
    """
    earliest = _earliest_assessment_by_substance(index)
    links = []
    for parent_uuid, metabolite_uuid in index.degradation_pairs:
        parent = index.substances.get(parent_uuid)
        metabolite = index.substances.get(metabolite_uuid)
        if not parent or not metabolite:
            continue
        parent_node_id = _node_id(parent)
        known_at = earliest.get(parent_node_id, fallback_known_at)
        links.append(
            DegradationLink(
                parent_substance_id=parent_node_id,
                metabolite_substance_id=_node_id(metabolite),
                source=SOURCE,
                known_at=known_at,
            )
        )
    return links


def assessments_from(index: OpenFoodToxIndex) -> list[SourceDocument]:
    """Dated EFSA scientific opinions, each linked to its assessed substance.

    ``known_at`` is the opinion's real evaluation date: this is what makes
    the temporal ``as_of`` cutoff meaningful rather than vacuous.
    """
    documents = []
    for row in index.assessments:
        subject = index.substances.get(row["subject_uuid"])
        if not subject:
            continue
        doi: str | None = row["doi"]
        doc_id = doi.removeprefix("doi:") if doi else f"efsa-dossier-{row['dossier_uuid'][:8]}"
        documents.append(
            SourceDocument(
                id=doc_id,
                title=row["title"],
                publisher=PUBLISHER,
                url=f"https://doi.org/{doi.removeprefix('doi:')}" if doi else None,
                published_at=row["published_at"],
                subject_substance_id=_node_id(subject),
                source=SOURCE,
                known_at=row["published_at"],
            )
        )
    return documents
